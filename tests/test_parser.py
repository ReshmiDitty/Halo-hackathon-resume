"""
Unit tests for document text extraction parser (parser.py).
"""

import io
import unittest
import docx
from parser import extract_text, _extract_docx, _extract_txt


class DummyUploadedFile:
    """Mock Streamlit UploadedFile class for testing."""

    def __init__(self, name: str, data: bytes):
        self.name = name
        self._data = data
        self._buffer = io.BytesIO(data)

    def read(self, *args, **kwargs):
        return self._buffer.read(*args, **kwargs)

    def seek(self, *args, **kwargs):
        return self._buffer.seek(*args, **kwargs)


class TestParser(unittest.TestCase):
    """Test suite for document parser module."""

    def test_extract_text_none_input(self):
        """Verify None input returns empty string without error."""
        self.assertEqual(extract_text(None), "")

    def test_extract_text_string_input(self):
        """Verify raw string input is returned stripped."""
        self.assertEqual(extract_text("  Sample Text  "), "Sample Text")

    def test_extract_txt_file(self):
        """Verify plain text file decoding."""
        content = "Software Engineer Resume\nPython, Docker, SQL"
        file_obj = DummyUploadedFile("resume.txt", content.encode("utf-8"))
        extracted = extract_text(file_obj)
        self.assertEqual(extracted, content)

    def test_extract_docx_file(self):
        """Verify docx document text extraction."""
        doc = docx.Document()
        doc.add_paragraph("Full Stack Developer")
        doc.add_paragraph("Skills: JavaScript, Python")

        doc_io = io.BytesIO()
        doc.save(doc_io)
        doc_io.seek(0)

        file_obj = DummyUploadedFile("resume.docx", doc_io.getvalue())
        extracted = extract_text(file_obj)
        self.assertIn("Full Stack Developer", extracted)
        self.assertIn("Skills: JavaScript, Python", extracted)

    def test_unsupported_file_extension(self):
        """Verify unsupported file extension raises ValueError."""
        file_obj = DummyUploadedFile("image.png", b"fake_png_data")
        with self.assertRaises(ValueError):
            extract_text(file_obj)


if __name__ == "__main__":
    unittest.main()
